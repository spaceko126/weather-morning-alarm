from __future__ import annotations

import datetime as dt
import html
import math
import os
import textwrap
from pathlib import Path
from typing import Iterable
from xml.sax.saxutils import escape

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from PIL import ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image as RLImage,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "project-brief"
DIAGRAMS = OUT / "diagrams"
RENDERED = OUT / "rendered"

DOCX_PATH = OUT / "WeatherMorningAlarm_Project_Brief.docx"
PDF_PATH = OUT / "WeatherMorningAlarm_Project_Brief.pdf"
DRAWIO_PATH = OUT / "weather-morning-architecture.drawio"

TODAY = dt.date(2026, 6, 25)
PROJECT = "Weather Morning Alarm"
VERSION = "0.1.2"
VERSION_CODE = "3"
APK_PATH = "release-apk/weather-morning-v3-0.1.2/weather-morning-v3-0.1.2.apk"
PROXY_URL = "https://<your-worker>.workers.dev"
APK_SHA256 = "B855AC996BB2102216730BC111013630FB7B96C29AB8B45320F908FE9D3F0BAE"

NAVY = "0B2545"
BLUE = "2E74B5"
SKY = "E8EEF5"
LIGHT = "F4F6F9"
MUTED = "5B6472"
GREEN = "1D7A46"
GOLD = "7A5A00"
RED = "9B1C1C"


def ensure_dirs() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    RENDERED.mkdir(parents=True, exist_ok=True)


def find_font() -> str:
    candidates = [
        Path(r"C:\Windows\Fonts\malgun.ttf"),
        Path(r"C:\Windows\Fonts\malgunsl.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return str(path)
    return ""


FONT_PATH = find_font()


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    bold_path = Path(r"C:\Windows\Fonts\malgunbd.ttf")
    try:
        if bold and bold_path.exists():
            return ImageFont.truetype(str(bold_path), size)
        if FONT_PATH:
            return ImageFont.truetype(FONT_PATH, size)
    except OSError:
        pass
    return ImageFont.load_default()


def hex_to_rgb(hex_color: str) -> tuple[int, int, int]:
    value = hex_color.strip("#")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def wrap_label(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [text]


def draw_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    subtitle: str = "",
    fill: str = "FFFFFF",
    outline: str = "C9D3E1",
) -> None:
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=14, fill=hex_to_rgb(fill), outline=hex_to_rgb(outline), width=2)
    title_font = pil_font(26, bold=True)
    body_font = pil_font(18)
    lines = wrap_label(draw, title, title_font, x2 - x1 - 36)
    y = y1 + 20
    for line in lines[:2]:
        draw.text((x1 + 18, y), line, fill=hex_to_rgb(NAVY), font=title_font)
        y += 32
    if subtitle:
        for line in wrap_label(draw, subtitle, body_font, x2 - x1 - 36)[:3]:
            draw.text((x1 + 18, y + 4), line, fill=hex_to_rgb(MUTED), font=body_font)
            y += 25


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = BLUE) -> None:
    sx, sy = start
    ex, ey = end
    draw.line((sx, sy, ex, ey), fill=hex_to_rgb(color), width=4)
    angle = math.atan2(ey - sy, ex - sx)
    size = 12
    points = [
        (ex, ey),
        (ex - size * math.cos(angle - math.pi / 6), ey - size * math.sin(angle - math.pi / 6)),
        (ex - size * math.cos(angle + math.pi / 6), ey - size * math.sin(angle + math.pi / 6)),
    ]
    draw.polygon(points, fill=hex_to_rgb(color))


def edge_point(
    box: tuple[int, int, int, int],
    toward: tuple[int, int],
    pad: int = 8,
) -> tuple[int, int]:
    x1, y1, x2, y2 = box
    cx = (x1 + x2) / 2
    cy = (y1 + y2) / 2
    tx, ty = toward
    dx = tx - cx
    dy = ty - cy
    if dx == 0 and dy == 0:
        return int(cx), int(cy)
    scale_x = ((x2 - x1) / 2 + pad) / abs(dx) if dx else float("inf")
    scale_y = ((y2 - y1) / 2 + pad) / abs(dy) if dy else float("inf")
    scale = min(scale_x, scale_y)
    return int(cx + dx * scale), int(cy + dy * scale)


def add_diagram_title(draw: ImageDraw.ImageDraw, title: str, subtitle: str) -> None:
    draw.text((50, 36), title, fill=hex_to_rgb(NAVY), font=pil_font(34, bold=True))
    draw.text((52, 82), subtitle, fill=hex_to_rgb(MUTED), font=pil_font(18))


def save_diagram(name: str, title: str, subtitle: str, boxes: list[dict], arrows: list[tuple[int, int]]) -> Path:
    img = PILImage.new("RGB", (1400, 900), hex_to_rgb("FFFFFF"))
    draw = ImageDraw.Draw(img)
    add_diagram_title(draw, title, subtitle)
    for item in boxes:
        draw_box(draw, item["box"], item["title"], item.get("subtitle", ""), item.get("fill", "FFFFFF"))
    centers = [((b["box"][0] + b["box"][2]) // 2, (b["box"][1] + b["box"][3]) // 2) for b in boxes]
    for start_idx, end_idx in arrows:
        start = edge_point(boxes[start_idx]["box"], centers[end_idx])
        end = edge_point(boxes[end_idx]["box"], centers[start_idx])
        draw_arrow(draw, start, end)
    path = DIAGRAMS / f"{name}.png"
    img.save(path)
    return path


def create_diagrams() -> dict[str, Path]:
    diagrams: dict[str, Path] = {}
    diagrams["overall"] = save_diagram(
        "01_overall_architecture",
        "Overall Architecture",
        "Galaxy-first Android app + Cloudflare weather proxy + local alarm engine",
        [
            {"box": (80, 170, 330, 290), "title": "User", "subtitle": "Alarm time, repeat days, location, weather rule", "fill": LIGHT},
            {"box": (430, 150, 730, 310), "title": "Android App", "subtitle": "Jetpack Compose UI, settings, history, alarm screen", "fill": SKY},
            {"box": (850, 140, 1180, 300), "title": "Cloudflare Worker", "subtitle": "Weather proxy, privacy page, cache, rate limit", "fill": LIGHT},
            {"box": (850, 420, 1180, 570), "title": "Weather Providers", "subtitle": "Open-Meteo default, KMA fallback for Korea", "fill": "F8FAFD"},
            {"box": (430, 430, 730, 610), "title": "Local Device Engine", "subtitle": "WorkManager precheck, exact alarm, decision store", "fill": "EFF7F1"},
            {"box": (120, 500, 330, 630), "title": "Alarm Output", "subtitle": "Full-screen weather alarm, sound, vibration, stop", "fill": "FFF7E8"},
        ],
        [(0, 1), (1, 2), (2, 3), (1, 4), (4, 5)],
    )
    diagrams["alarm"] = save_diagram(
        "02_alarm_decision_flow",
        "Alarm Decision Flow",
        "Weather is checked before the alarm. Ring only when the saved decision says yes.",
        [
            {"box": (70, 160, 310, 285), "title": "Settings saved", "subtitle": "Time, repeat, location, POP %, optional rain amount"},
            {"box": (390, 160, 650, 285), "title": "Schedule", "subtitle": "Precheck + exact alarm"},
            {"box": (730, 160, 1030, 285), "title": "Precheck worker", "subtitle": "Fetch target forecast rows"},
            {"box": (1080, 160, 1320, 285), "title": "Decision engine", "subtitle": "Rain/snow/POP/mm thresholds"},
            {"box": (730, 470, 1030, 610), "title": "Saved decision", "subtitle": "Ring/skip reason + forecast rows"},
            {"box": (390, 470, 650, 610), "title": "Exact alarm fires", "subtitle": "Read saved decision"},
            {"box": (70, 470, 310, 610), "title": "Ring or skip", "subtitle": "Show alarm screen or save skipped history"},
        ],
        [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (5, 6)],
    )
    diagrams["data"] = save_diagram(
        "03_data_storage_flow",
        "Data and Storage Flow",
        "User settings and weather decisions stay local. Backend hides shared provider keys.",
        [
            {"box": (80, 155, 350, 285), "title": "Settings Store", "subtitle": "DataStore: alarm, sound, vibration, location"},
            {"box": (520, 155, 820, 285), "title": "Decision Store", "subtitle": "JSON: last decision + forecast snapshots"},
            {"box": (990, 155, 1290, 285), "title": "History DB", "subtitle": "Room: max-visible history, reason, weather"},
            {"box": (80, 470, 350, 610), "title": "Backend Proxy", "subtitle": "No raw API key in Android app"},
            {"box": (520, 470, 820, 610), "title": "Privacy Route", "subtitle": "Public /privacy for distribution"},
            {"box": (990, 470, 1290, 610), "title": "Logs", "subtitle": "Structured, no raw coordinates/API keys"},
        ],
        [(0, 1), (1, 2), (3, 4), (3, 5)],
    )
    diagrams["code"] = save_diagram(
        "04_code_structure",
        "Code Structure",
        "Layered Android code plus a small Worker backend and release automation scripts.",
        [
            {"box": (70, 150, 350, 290), "title": "presentation/", "subtitle": "Compose screens, alarm UI, settings, history"},
            {"box": (400, 150, 680, 290), "title": "domain/", "subtitle": "Models, weather decision use cases"},
            {"box": (730, 150, 1010, 290), "title": "data/", "subtitle": "Local stores, Room DB, weather repositories"},
            {"box": (1060, 150, 1340, 290), "title": "alarm + worker/", "subtitle": "Receivers, scheduler, WorkManager"},
            {"box": (220, 500, 520, 640), "title": "backend/weather-proxy", "subtitle": "Cloudflare Worker weather API proxy"},
            {"box": (620, 500, 920, 640), "title": "scripts/", "subtitle": "Build, QA, APK packaging, smoke checks"},
            {"box": (1010, 500, 1290, 640), "title": "docs/", "subtitle": "Progress, release, QA, handoff docs"},
        ],
        [(0, 1), (1, 2), (2, 3), (4, 5), (5, 6)],
    )
    diagrams["release"] = save_diagram(
        "05_release_pipeline",
        "Release Pipeline",
        "Debug validation, proxy smoke, signed APK build, verification, package folder.",
        [
            {"box": (80, 160, 330, 290), "title": "Unit + build", "subtitle": "scripts/test-debug.ps1"},
            {"box": (410, 160, 660, 290), "title": "Device QA", "subtitle": "Galaxy adb tests and alarm signal paths"},
            {"box": (740, 160, 990, 290), "title": "Proxy smoke", "subtitle": "Production Worker + privacy contact"},
            {"box": (1070, 160, 1320, 290), "title": "Signed APK", "subtitle": "v3 / 0.1.2"},
            {"box": (410, 500, 660, 630), "title": "Verify APK", "subtitle": "signature, version, URL, SHA-256"},
            {"box": (740, 500, 990, 630), "title": "Package", "subtitle": "APK + README + privacy + notes"},
        ],
        [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5)],
    )
    diagrams["qa"] = save_diagram(
        "06_qa_coverage",
        "QA Coverage",
        "Strong debug/release checks exist; a few real-world device cases remain.",
        [
            {"box": (90, 160, 390, 310), "title": "Verified", "subtitle": "Unit tests, backend tests, debug alarm signals, release APK verification", "fill": "EFF7F1"},
            {"box": (550, 160, 850, 310), "title": "Partially verified", "subtitle": "adb confirms signal/UI/stop; human audio/haptic is subjective", "fill": "FFF7E8"},
            {"box": (1010, 160, 1310, 310), "title": "Remaining", "subtitle": "Real scheduled release alarm, lock screen, DND, battery modes", "fill": "FBEFEF"},
            {"box": (320, 520, 610, 660), "title": "Safe to share APK", "subtitle": "First-time sideload path is ready"},
            {"box": (790, 520, 1080, 660), "title": "Same phone caveat", "subtitle": "Debug app must be uninstalled before release install"},
        ],
        [(0, 1), (1, 2), (3, 4)],
    )
    return diagrams


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str = NAVY, size: int = 10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(*hex_to_rgb(color))
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")


def set_docx_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*hex_to_rgb(color))
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25


def add_docx_para(doc: Document, text: str, size: int = 11, bold: bool = False, color: str = "000000") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(*hex_to_rgb(color))
    run.bold = bold


def add_docx_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.add_run(item)


def add_docx_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, (label, value) in enumerate(rows):
        cells = table.rows[row_idx].cells
        cells[0].width = Inches(1.181)
        cells[1].width = Inches(5.319)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cells[0], SKY)
        set_cell_text(cells[0], label, bold=True, color=NAVY, size=9)
        set_cell_text(cells[1], value, color="1F2937", size=9)


def build_docx(diagrams: dict[str, Path]) -> None:
    doc = Document()
    set_docx_styles(doc)

    section = doc.sections[0]
    section.header.paragraphs[0].text = f"{PROJECT} | Project Brief"
    section.footer.paragraphs[0].text = f"Generated {TODAY.isoformat()} | Page"

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run(PROJECT)
    run.bold = True
    run.font.size = Pt(25)
    run.font.color.rgb = RGBColor(*hex_to_rgb(NAVY))
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    add_docx_para(
        doc,
        "비·눈이 예상되는 아침에만 사용자를 더 일찍 깨우는 Galaxy-first 조건부 날씨 알람 앱.",
        size=12,
        color=MUTED,
    )

    add_docx_table(
        doc,
        [
            ("상태", "Android MVP 구현 완료, Cloudflare Worker 배포/스모크 통과, signed APK v0.1.2 준비"),
            ("핵심 가치", "비나 눈이 오는 출근 전 시간대에만 알람을 울려 불필요한 알람을 줄임"),
            ("릴리스 파일", APK_PATH),
            ("백엔드", PROXY_URL),
            ("검증", "debug/unit/backend/device signal/release APK verification 통과. 실제 장기 반복·DND·배터리 모드 QA는 남음"),
        ],
    )

    doc.add_heading("1. 한 페이지 요약", level=1)
    add_docx_bullets(
        doc,
        [
            "사용자는 알람 시간, 반복 요일, 위치, 강수확률, 선택 강수량 기준을 설정한다.",
            "앱은 알람 전 지정 범위의 날씨를 확인하고, 조건이 맞을 때만 정확 알람을 울린다.",
            "날씨 API 키는 Android 앱에 넣지 않고 Cloudflare Worker가 프록시한다.",
            "알람 화면은 오늘 상태, 현재/시간별 예보, 큰 중지 버튼을 보여준다.",
            "현재 배포 형태는 APK 직접 설치이며 Play Console 출시는 계정 작업 후 가능하다.",
        ],
    )

    doc.add_page_break()
    for title_text, key, caption in [
        ("2. 전체 아키텍처", "overall", "앱, 로컬 알람 엔진, Cloudflare Worker, 외부 날씨 제공자의 관계"),
        ("3. 알람 판단 흐름", "alarm", "사전 예보 확인부터 알람 울림/건너뜀까지의 런타임 흐름"),
        ("4. 데이터와 저장소", "data", "로컬 저장 데이터와 백엔드가 감추는 외부 API 키 경계"),
        ("5. 코드 구조", "code", "Android 레이어, Worker, 스크립트, 문서 구조"),
        ("6. 릴리스 파이프라인", "release", "검증, 서명, APK 패키징 순서"),
        ("7. QA 범위", "qa", "검증 완료 영역과 아직 실제 환경 확인이 필요한 영역"),
    ]:
        doc.add_heading(title_text, level=1)
        add_docx_para(doc, caption, color=MUTED)
        doc.add_picture(str(diagrams[key]), width=Inches(6.45))
        if key != "qa":
            doc.add_page_break()

    doc.add_heading("8. 다음 작업", level=1)
    add_docx_table(
        doc,
        [
            ("우선", "릴리스 APK를 새 기기 또는 debug 앱 삭제 후 설치해 실제 예약 알람 경로 확인"),
            ("QA", "잠금화면, DND/무음, 배터리 최적화, 여러 날짜 반복 알람 확인"),
            ("배포", "Play Console 계정 결제/앱 등록 후 AAB 업로드, 개인정보/권한 문구 입력"),
            ("고도화", "전체 일일 날씨 화면, 더 많은 지역/국가별 제공자, iOS 검토"),
        ],
    )
    doc.save(DOCX_PATH)


def register_pdf_font() -> str:
    if FONT_PATH:
        pdfmetrics.registerFont(TTFont("Malgun", FONT_PATH))
        return "Malgun"
    return "Helvetica"


def pdf_styles(font_name: str):
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="KTitle",
            parent=styles["Title"],
            fontName=font_name,
            fontSize=24,
            leading=30,
            textColor=colors.HexColor(f"#{NAVY}"),
            alignment=TA_LEFT,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="KSubtitle",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=11,
            leading=16,
            textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=10,
        )
    )
    styles.add(
        ParagraphStyle(
            name="KHeading",
            parent=styles["Heading1"],
            fontName=font_name,
            fontSize=16,
            leading=22,
            textColor=colors.HexColor(f"#{BLUE}"),
            spaceBefore=8,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="KBody",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=10.5,
            leading=15,
            textColor=colors.HexColor("#111827"),
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="KSmall",
            parent=styles["BodyText"],
            fontName=font_name,
            fontSize=9,
            leading=13,
            textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=4,
        )
    )
    return styles


def pdf_table(rows: list[tuple[str, str]], font_name: str) -> Table:
    data = [[Paragraph(f"<b>{html.escape(a)}</b>", ParagraphStyle("L", fontName=font_name, fontSize=8.5, leading=11)),
             Paragraph(html.escape(b), ParagraphStyle("V", fontName=font_name, fontSize=8.5, leading=11))]
            for a, b in rows]
    table = Table(data, colWidths=[1.18 * inch, 5.12 * inch], hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D3E1")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(f"#{SKY}")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def build_pdf(diagrams: dict[str, Path]) -> None:
    font = register_pdf_font()
    styles = pdf_styles(font)
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.65 * inch,
        title=f"{PROJECT} Project Brief",
    )
    story = [
        Paragraph(PROJECT, styles["KTitle"]),
        Paragraph("비·눈이 예상되는 아침에만 사용자를 더 일찍 깨우는 Galaxy-first 조건부 날씨 알람 앱.", styles["KSubtitle"]),
        pdf_table(
            [
                ("현재 상태", "Android MVP 구현 완료, Worker 배포/스모크 통과, signed APK v0.1.2 준비"),
                ("사용자 가치", "비·눈이 오는 출근 전 시간대에만 알람을 울려 불필요한 알람을 줄임"),
                ("릴리스 파일", APK_PATH),
                ("백엔드", PROXY_URL),
                ("APK SHA-256", APK_SHA256),
            ],
            font,
        ),
        Spacer(1, 10),
        Paragraph("핵심 설계", styles["KHeading"]),
    ]
    for item in [
        "설정: 알람 시간, 반복 요일, 위치, 강수확률, 선택 강수량 기준",
        "사전 확인: WorkManager/정확 알람으로 예보를 먼저 확인하고 판단 결과를 저장",
        "울림 조건: 비·눈·강수확률·선택 mm 기준이 맞을 때만 전체화면 알람",
        "보안: 날씨 API 키는 앱에 넣지 않고 Cloudflare Worker에서 관리",
        "배포: APK 직접 설치 가능, Play Console 출시는 계정/스토어 절차 필요",
    ]:
        story.append(Paragraph(f"• {html.escape(item)}", styles["KBody"]))

    story.append(PageBreak())
    for heading, key, note in [
        ("전체 아키텍처", "overall", "앱, 로컬 알람 엔진, Cloudflare Worker, 외부 날씨 제공자의 관계"),
        ("알람 판단 흐름", "alarm", "사전 예보 확인부터 알람 울림/건너뜀까지의 런타임 흐름"),
        ("데이터와 저장소", "data", "로컬 저장 데이터와 백엔드가 감추는 외부 API 키 경계"),
        ("코드 구조", "code", "Android 레이어, Worker, 스크립트, 문서 구조"),
        ("릴리스 파이프라인", "release", "검증, 서명, APK 패키징 순서"),
        ("QA 범위", "qa", "검증 완료 영역과 아직 실제 환경 확인이 필요한 영역"),
    ]:
        story.extend(
            [
                Paragraph(heading, styles["KHeading"]),
                Paragraph(note, styles["KSmall"]),
                RLImage(str(diagrams[key]), width=7.15 * inch, height=4.60 * inch),
                PageBreak(),
            ]
        )
    story.pop()
    story.extend(
        [
            Spacer(1, 10),
            Paragraph("남은 작업", styles["KHeading"]),
            pdf_table(
                [
                    ("실기기 QA", "릴리스 APK 설치 후 실제 예약 알람, 잠금화면, DND/무음, 배터리 최적화 확인"),
                    ("스토어", "Play Console 계정 결제/앱 등록/AAB 업로드/권한 문구 입력"),
                    ("고도화", "전체 일일 날씨, 더 많은 제공자, iOS 검토"),
                ],
                font,
            ),
        ]
    )
    doc.build(story)


def mx_cell(id_: str, value: str, x: int, y: int, w: int, h: int, fill: str = SKY) -> str:
    return (
        f'<mxCell id="{id_}" value="{escape(value)}" '
        f'style="rounded=1;whiteSpace=wrap;html=1;fillColor=#{fill};strokeColor=#C9D3E1;fontColor=#{NAVY};fontStyle=1;" '
        f'vertex="1" parent="1"><mxGeometry x="{x}" y="{y}" width="{w}" height="{h}" as="geometry"/></mxCell>'
    )


def mx_edge(id_: str, source: str, target: str) -> str:
    return (
        f'<mxCell id="{id_}" style="edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;'
        f'jettySize=auto;html=1;strokeColor=#{BLUE};endArrow=block;endFill=1;" edge="1" parent="1" '
        f'source="{source}" target="{target}"><mxGeometry relative="1" as="geometry"/></mxCell>'
    )


def graph_page(name: str, cells: list[str]) -> str:
    inner = (
        '<mxGraphModel dx="1200" dy="800" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" '
        'arrows="1" fold="1" page="1" pageScale="1" pageWidth="1400" pageHeight="900" math="0" shadow="0">'
        '<root><mxCell id="0"/><mxCell id="1" parent="0"/>'
        + "".join(cells)
        + "</root></mxGraphModel>"
    )
    return f'<diagram id="{escape(name.lower().replace(" ", "-"))}" name="{escape(name)}"><![CDATA[{inner}]]></diagram>'


def build_drawio() -> None:
    pages = [
        graph_page(
            "01 Overall Architecture",
            [
                mx_cell("u", "User<br>settings + location", 60, 140, 220, 90, LIGHT),
                mx_cell("a", "Android App<br>Compose UI + local alarm", 380, 120, 260, 120),
                mx_cell("w", "Cloudflare Worker<br>weather proxy + privacy", 760, 120, 290, 120, LIGHT),
                mx_cell("p", "Weather Providers<br>Open-Meteo + KMA fallback", 760, 360, 290, 120, "F8FAFD"),
                mx_cell("e", "Alarm Engine<br>precheck + exact alarm", 380, 360, 260, 120, "EFF7F1"),
                mx_cell("o", "Alarm Output<br>weather screen + sound/vibration", 60, 360, 220, 120, "FFF7E8"),
                mx_edge("e1", "u", "a"),
                mx_edge("e2", "a", "w"),
                mx_edge("e3", "w", "p"),
                mx_edge("e4", "a", "e"),
                mx_edge("e5", "e", "o"),
            ],
        ),
        graph_page(
            "02 Alarm Flow",
            [
                mx_cell("s", "Settings saved", 60, 160, 190, 80),
                mx_cell("sc", "Schedule precheck + exact alarm", 330, 160, 230, 80),
                mx_cell("wk", "WeatherCheckWorker fetches forecast", 650, 160, 250, 80),
                mx_cell("de", "Decision engine applies weather rule", 980, 160, 250, 80),
                mx_cell("ds", "Decision store", 650, 390, 250, 80, "EFF7F1"),
                mx_cell("ar", "AlarmReceiver", 330, 390, 230, 80, "FFF7E8"),
                mx_cell("rs", "Ring or skip + history", 60, 390, 190, 80, "FFF7E8"),
                mx_edge("a1", "s", "sc"),
                mx_edge("a2", "sc", "wk"),
                mx_edge("a3", "wk", "de"),
                mx_edge("a4", "de", "ds"),
                mx_edge("a5", "ds", "ar"),
                mx_edge("a6", "ar", "rs"),
            ],
        ),
        graph_page(
            "03 Data Storage",
            [
                mx_cell("st", "DataStore<br>settings, sound, vibration", 80, 170, 270, 100),
                mx_cell("ld", "Local decision JSON<br>forecast snapshots", 520, 170, 270, 100),
                mx_cell("hi", "Room history<br>reason + weather", 960, 170, 270, 100),
                mx_cell("pr", "Worker cache/rate limit<br>provider key hidden", 300, 430, 300, 100, LIGHT),
                mx_cell("pv", "Public privacy route<br>/privacy", 760, 430, 300, 100, LIGHT),
                mx_edge("d1", "st", "ld"),
                mx_edge("d2", "ld", "hi"),
                mx_edge("d3", "pr", "pv"),
            ],
        ),
        graph_page(
            "04 Code Structure",
            [
                mx_cell("pre", "presentation/<br>Compose screens", 70, 150, 230, 100),
                mx_cell("dom", "domain/<br>models + use cases", 370, 150, 230, 100),
                mx_cell("dat", "data/<br>local + repositories", 670, 150, 230, 100),
                mx_cell("alw", "alarm/ + worker/<br>receivers + WorkManager", 970, 150, 260, 100),
                mx_cell("be", "backend/weather-proxy<br>Cloudflare Worker", 220, 440, 270, 100, LIGHT),
                mx_cell("sc", "scripts/<br>QA + release automation", 570, 440, 270, 100, LIGHT),
                mx_cell("do", "docs/<br>progress + handoff", 920, 440, 270, 100, LIGHT),
                mx_edge("c1", "pre", "dom"),
                mx_edge("c2", "dom", "dat"),
                mx_edge("c3", "dat", "alw"),
                mx_edge("c4", "be", "sc"),
                mx_edge("c5", "sc", "do"),
            ],
        ),
        graph_page(
            "05 Release Pipeline",
            [
                mx_cell("t", "test-debug.ps1", 80, 170, 220, 80),
                mx_cell("d", "device/signal QA", 380, 170, 220, 80),
                mx_cell("p", "proxy smoke", 680, 170, 220, 80),
                mx_cell("b", "signed APK build", 980, 170, 220, 80),
                mx_cell("v", "verify APK", 380, 410, 220, 80, "EFF7F1"),
                mx_cell("pkg", "release package", 680, 410, 220, 80, "EFF7F1"),
                mx_edge("r1", "t", "d"),
                mx_edge("r2", "d", "p"),
                mx_edge("r3", "p", "b"),
                mx_edge("r4", "b", "v"),
                mx_edge("r5", "v", "pkg"),
            ],
        ),
    ]
    DRAWIO_PATH.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<mxfile host="app.diagrams.net" modified="2026-06-25T00:00:00.000Z" agent="Codex" version="24.7.17">'
        + "".join(pages)
        + "</mxfile>",
        encoding="utf-8",
    )


def render_pdf_pages() -> None:
    doc = fitz.open(str(PDF_PATH))
    for idx, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(1.35, 1.35), alpha=False)
        pix.save(str(RENDERED / f"pdf_page_{idx + 1:02d}.png"))
    doc.close()


def main() -> None:
    ensure_dirs()
    diagrams = create_diagrams()
    build_docx(diagrams)
    build_pdf(diagrams)
    build_drawio()
    render_pdf_pages()
    print(f"DOCX={DOCX_PATH}")
    print(f"PDF={PDF_PATH}")
    print(f"DRAWIO={DRAWIO_PATH}")
    print(f"RENDERED={RENDERED}")


if __name__ == "__main__":
    main()
